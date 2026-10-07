import os
import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image as RLImage, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(r"C:\Users\sylva\OneDrive\Desktop\DVM Internship\Timetable")
TEMPLATE = Path(r"C:\Users\sylva\OneDrive\Desktop\DVM Internship\UCVM Staging Environment Doc.docx")
OUT = ROOT / "docs" / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)
DOCX_OUT = OUT / "UCVM Staging Environment Guide.docx"
PDF_OUT = OUT / "UCVM Staging Environment Guide.pdf"
DIAGRAM = OUT / "staging-architecture.png"
TODAY = "October 6, 2026"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        elem = borders.find(tag)
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_text(paragraph, text, bold=False, size=None, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def add_para(doc, text="", style=None, before=0, after=6, bold_lead=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True)
        add_text(p, text[len(bold_lead):])
    else:
        add_text(p, text)
    return p


def add_bullets(doc, items):
    for item in items:
        # The supplied template does not define Word's built-in List Bullet
        # style, so create a stable, template-compatible bullet paragraph.
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.paragraph_format.space_after = Pt(2)
        add_text(p, "• " + item)


def add_docx_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "2E74B5")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, header, bold=True, size=9, color=(255, 255, 255))
        if widths:
            cell.width = Inches(widths[index])
    for r, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            if r % 2:
                set_cell_shading(cell, "F3F7FB")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, value, size=9)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def font(size, bold=False):
    for candidate in ([r"C:\Windows\Fonts\arialbd.ttf"] if bold else []) + [r"C:\Windows\Fonts\arial.ttf"]:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def diagram():
    image = Image.new("RGB", (1800, 720), "white")
    draw = ImageDraw.Draw(image)
    blue, pale, dark, green = "#2E74B5", "#EAF3FB", "#1F1F1F", "#70AD47"
    boxes = [
        (55, 160, 505, 510, "Production", "GitHub Pages\nFirestore: timetable-23438\nProduction users and data"),
        (675, 160, 1125, 510, "Controlled refresh", "Authorized operator\nRead-only production key\nVerified copy and rollback"),
        (1295, 160, 1745, 510, "Staging", "Firebase Hosting\nFirestore: ucvm-timetable-staging\nSeparate Auth and test data"),
    ]
    for x1, y1, x2, y2, title, body in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=pale, outline=blue, width=5)
        draw.rounded_rectangle((x1, y1, x2, y1 + 85), radius=22, fill=blue, outline=blue, width=2)
        draw.text((x1 + 28, y1 + 23), title, font=font(34, True), fill="white")
        draw.multiline_text((x1 + 28, y1 + 130), body, font=font(28), fill=dark, spacing=15)
    draw.line((510, 335, 665, 335), fill=green, width=9)
    draw.polygon([(665, 335), (630, 315), (630, 355)], fill=green)
    draw.line((1130, 335, 1285, 335), fill=green, width=9)
    draw.polygon([(1285, 335), (1250, 315), (1250, 355)], fill=green)
    draw.text((595, 82), "Selected timetable data", font=font(22, True), fill=dark)
    draw.text((1175, 82), "Fresh staging dataset", font=font(22, True), fill=dark)
    image.save(DIAGRAM)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def build_docx():
    shutil.copy2(TEMPLATE, DOCX_OUT)
    doc = Document(DOCX_OUT)
    clear_body(doc)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    heading = doc.styles["Heading 1"]
    heading.font.name = "Calibri"
    heading.font.size = Pt(15)
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(15)
    heading.paragraph_format.space_after = Pt(5)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("UCVM Staging Environment Guide")
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph(style="No Spacing")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, f"Updated: {TODAY}")
    p = doc.add_paragraph(style="No Spacing")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "Authors: Jasneet Singh & Eze Adiele")
    doc.add_paragraph()

    doc.add_paragraph("Purpose and scope", style="Heading 1")
    add_para(doc, "The UCVM staging environment is a separate, safe copy of the timetable system used to test changes before they affect the live timetable. It has its own website, Firestore database, authentication accounts, and operational credentials. Production remains independent.")
    add_docx_table(doc, ["Production", "Staging"], [
        ("GitHub Pages site and Firebase project timetable-23438", "Firebase Hosting site and Firebase project ucvm-timetable-staging"),
        ("Live timetable and live user activity", "Testing data and selected staging testers"),
        ("Never written by the staging refresh tool", "Can be refreshed from selected production timetable data"),
    ], [3.2, 3.2])

    doc.add_paragraph("How the environment works", style="Heading 1")
    add_para(doc, "Staging is intentionally isolated. A page served from Firebase Hosting loads the staging configuration, connects only to the staging Firebase project, and keeps its browser cache separate from production.")
    doc.add_picture(str(DIAGRAM), width=Inches(6.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "The refresh process is deliberately one-way: it reads only selected timetable data from production and writes only to staging. Production users, authentication records, change requests, access requests, and audit logs are not copied.")

    doc.add_paragraph("What is refreshed", style="Heading 1")
    add_docx_table(doc, ["Included", "Not copied from production"], [
        ("sessions", "Firebase Authentication users and passwords"),
        ("roster", "users and authorized user profiles"),
        ("settings/rosterNotice", "change requests, access requests, history, change logs, audit logs"),
        ("A new staging-only sessionsVersion timestamp", "Production sessionsVersion timestamp"),
    ], [3.2, 3.2])
    add_para(doc, "During a refresh, staging displays a short maintenance state. After the copied data is verified, the staging sessionsVersion timestamp is updated once. This tells open staging browsers to replace any older cached timetable data with the refreshed staging data.")

    doc.add_paragraph("Where to find each component", style="Heading 1")
    add_docx_table(doc, ["Component", "Where to find it", "Why it matters"], [
        ("Staging website", "https://ucvm-timetable-staging.web.app", "Safe place to test the deployed application."),
        ("Source code", "Timetable repository, staging-branch", "Contains the staging application and operations tools."),
        ("Hosting settings", "firebase.json and .firebaserc", "Targets the staging Firebase project and excludes the Excel data feed."),
        ("Firestore rules", "firestore.rules", "Controls browser access and blocks browser writes during a refresh."),
        ("Refresh tool", "scripts/refresh-staging.mjs", "Copies approved data, verifies it, and rolls back if necessary."),
        ("Tester allowlist", "scripts/seed-staging-allowlist.mjs", "Adds selected testers without copying production accounts."),
        ("Operations guide", "docs/staging-environment-runbook.md", "Exact setup, deployment, refresh, and recovery commands."),
    ], [1.25, 2.25, 2.9])

    doc.add_paragraph("Firebase Hosting and Google Cloud access", style="Heading 1")
    add_para(doc, "Firebase Hosting publishes the staging branch to the staging site. The site configuration deliberately does not publish data-feed.html, documentation, scripts, dependencies, or credential files.")
    add_para(doc, "Two purpose-limited Google Cloud service accounts support a refresh. Their exact email addresses and JSON keys are not kept in the repository. Find and manage them in Google Cloud Console by selecting the appropriate project, then opening IAM & Admin > Service Accounts.")
    add_docx_table(doc, ["Access role", "Project", "Permitted use"], [
        ("Production reader", "timetable-23438", "Read Firestore data required for the approved staging copy. It cannot write production data."),
        ("Staging writer", "ucvm-timetable-staging", "Write, verify, and if needed roll back staging data. It cannot access production."),
    ], [1.55, 1.75, 3.1])
    add_para(doc, "Keys must be stored outside the repository and outside OneDrive. The refresh tool checks both project IDs before it begins. Browser users never receive these keys.")

    doc.add_paragraph("Access and day to day operation", style="Heading 1")
    add_bullets(doc, [
        "Staging users may use the same email address as production, but they need a separate staging password because Firebase Authentication is a separate user pool.",
        "A staging administrator creates selected tester entries in the staging allowlist. Users then create or reset their own staging credentials through the site.",
        "Before a refresh, run the read-only dry run. Apply a refresh only from a trusted operator computer with both service-account paths set for that PowerShell session.",
        "If a refresh fails, the tool restores the previous staging snapshot and leaves the cache timestamp unchanged. If rollback fails, staging stays in maintenance until an authorized operator resolves it.",
    ])

    doc.add_paragraph("Common questions", style="Heading 1")
    faq = [
        ("Will staging change the live timetable?", "No. The refresh tool reads production but never writes to it. Staging and production use different Firebase projects."),
        ("Why do testers need a separate password?", "Authentication is intentionally separate. This keeps staging tests, user IDs, and password resets away from production."),
        ("Why can a refresh affect an open browser?", "Browsers cache timetable data for speed. The final staging-only version update tells them to fetch the refreshed data."),
        ("Can the staging Excel data feed be used?", "No. data-feed.html is intentionally excluded from Firebase Hosting to avoid an unnecessary export endpoint and extra reads in staging."),
        ("Where are credentials stored?", "Service-account keys are held by authorized operators outside the repository. Passwords and JSON keys should never be added to source control or shared in this guide."),
        ("Who should run a refresh?", "Only a trusted technical operator with the two purpose-limited Google Cloud service-account keys and the staging runbook."),
    ]
    for question, answer in faq:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        add_text(p, question, bold=True)
        add_para(doc, answer, after=4)

    doc.core_properties.title = "UCVM Staging Environment Guide"
    doc.core_properties.subject = "UCVM timetable staging environment"
    doc.core_properties.author = "Sylva Nax"  # Preserves template document metadata.
    doc.core_properties.last_modified_by = "Sylva Nax"
    doc.save(DOCX_OUT)


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("Title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=23, leading=28, alignment=TA_CENTER, textColor=colors.black, spaceAfter=7),
        "meta": ParagraphStyle("Meta", parent=styles["Normal"], fontName="Helvetica", fontSize=10, leading=13, alignment=TA_CENTER, spaceAfter=1),
        "h1": ParagraphStyle("Heading", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=colors.black, spaceBefore=15, spaceAfter=6),
        "body": ParagraphStyle("Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=14, textColor=colors.black, spaceAfter=7),
        "bullet": ParagraphStyle("Bullet", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=13, leftIndent=16, firstLineIndent=-8, spaceAfter=3),
        "question": ParagraphStyle("Question", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=4, spaceAfter=1),
    }


def p(text, style):
    return Paragraph(text.replace("&", "&amp;"), style)


def table_pdf(headers, rows, widths):
    data = [[p(h, PDF_STYLES["body"]) for h in headers]] + [[p(cell, PDF_STYLES["body"]) for cell in row] for row in rows]
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E74B5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]
    for index in range(1, len(data)):
        if index % 2 == 0:
            style.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#F3F7FB")))
    table.setStyle(TableStyle(style))
    return table


def build_pdf():
    global PDF_STYLES
    PDF_STYLES = pdf_styles()
    doc = SimpleDocTemplate(str(PDF_OUT), pagesize=letter, rightMargin=0.72*inch, leftMargin=0.72*inch, topMargin=0.67*inch, bottomMargin=0.62*inch, title="UCVM Staging Environment Guide", author="Sylva Nax")
    story = [
        p("UCVM Staging Environment Guide", PDF_STYLES["title"]),
        p(f"Updated: {TODAY}", PDF_STYLES["meta"]),
        p("Authors: Jasneet Singh & Eze Adiele", PDF_STYLES["meta"]), Spacer(1, 13),
        p("Purpose and scope", PDF_STYLES["h1"]),
        p("The UCVM staging environment is a separate, safe copy of the timetable system used to test changes before they affect the live timetable. It has its own website, Firestore database, authentication accounts, and operational credentials. Production remains independent.", PDF_STYLES["body"]),
        table_pdf(["Production", "Staging"], [
            ("GitHub Pages site and Firebase project timetable-23438", "Firebase Hosting site and Firebase project ucvm-timetable-staging"),
            ("Live timetable and live user activity", "Testing data and selected staging testers"),
            ("Never written by the staging refresh tool", "Can be refreshed from selected production timetable data"),
        ], [3.2*inch, 3.2*inch]),
        p("How the environment works", PDF_STYLES["h1"]),
        p("Staging is intentionally isolated. A page served from Firebase Hosting loads the staging configuration, connects only to the staging Firebase project, and keeps its browser cache separate from production.", PDF_STYLES["body"]),
        RLImage(str(DIAGRAM), width=6.35*inch, height=2.54*inch),
        p("The refresh process reads only selected timetable data from production and writes only to staging. Production users, authentication records, change requests, access requests, and audit logs are not copied.", PDF_STYLES["body"]),
        p("What is refreshed", PDF_STYLES["h1"]),
        table_pdf(["Included", "Not copied from production"], [
            ("sessions", "Firebase Authentication users and passwords"), ("roster", "users and authorized user profiles"),
            ("settings/rosterNotice", "change requests, access requests, history, change logs, audit logs"),
            ("A new staging-only sessionsVersion timestamp", "Production sessionsVersion timestamp"),
        ], [3.2*inch, 3.2*inch]),
        p("During a refresh, staging displays a short maintenance state. After the copied data is verified, the staging sessionsVersion timestamp is updated once. This tells open staging browsers to replace any older cached timetable data with the refreshed staging data.", PDF_STYLES["body"]),
        p("Where to find each component", PDF_STYLES["h1"]),
        table_pdf(["Component", "Where to find it", "Why it matters"], [
            ("Staging website", "https://ucvm-timetable-staging.web.app", "Safe place to test the deployed application."),
            ("Source code", "Timetable repository, staging-branch", "Contains the staging application and operations tools."),
            ("Hosting settings", "firebase.json and .firebaserc", "Targets the staging Firebase project and excludes the Excel data feed."),
            ("Firestore rules", "firestore.rules", "Controls browser access and blocks browser writes during a refresh."),
            ("Refresh tool", "scripts/refresh-staging.mjs", "Copies approved data, verifies it, and rolls back if necessary."),
            ("Tester allowlist", "scripts/seed-staging-allowlist.mjs", "Adds selected testers without copying production accounts."),
            ("Operations guide", "docs/staging-environment-runbook.md", "Exact setup, deployment, refresh, and recovery commands."),
        ], [1.2*inch, 2.3*inch, 2.9*inch]),
        p("Firebase Hosting and Google Cloud access", PDF_STYLES["h1"]),
        p("Firebase Hosting publishes the staging branch to the staging site. The site configuration deliberately does not publish data-feed.html, documentation, scripts, dependencies, or credential files.", PDF_STYLES["body"]),
        p("Two purpose-limited Google Cloud service accounts support a refresh. Their exact email addresses and JSON keys are not kept in the repository. Find and manage them in Google Cloud Console by selecting the appropriate project, then opening IAM & Admin > Service Accounts.", PDF_STYLES["body"]),
        table_pdf(["Access role", "Project", "Permitted use"], [
            ("Production reader", "timetable-23438", "Read Firestore data required for the approved staging copy. It cannot write production data."),
            ("Staging writer", "ucvm-timetable-staging", "Write, verify, and if needed roll back staging data. It cannot access production."),
        ], [1.55*inch, 1.75*inch, 3.1*inch]),
        p("Keys must be stored outside the repository and outside OneDrive. The refresh tool checks both project IDs before it begins. Browser users never receive these keys.", PDF_STYLES["body"]),
        p("Access and day to day operation", PDF_STYLES["h1"]),
    ]
    for item in [
        "Staging users may use the same email address as production, but they need a separate staging password because Firebase Authentication is a separate user pool.",
        "A staging administrator creates selected tester entries in the staging allowlist. Users then create or reset their own staging credentials through the site.",
        "Before a refresh, run the read-only dry run. Apply a refresh only from a trusted operator computer with both service-account paths set for that PowerShell session.",
        "If a refresh fails, the tool restores the previous staging snapshot and leaves the cache timestamp unchanged. If rollback fails, staging stays in maintenance until an authorized operator resolves it.",
    ]:
        story.append(p("- " + item, PDF_STYLES["bullet"]))
    story.append(p("Common questions", PDF_STYLES["h1"]))
    for q, a in [
        ("Will staging change the live timetable?", "No. The refresh tool reads production but never writes to it. Staging and production use different Firebase projects."),
        ("Why do testers need a separate password?", "Authentication is intentionally separate. This keeps staging tests, user IDs, and password resets away from production."),
        ("Why can a refresh affect an open browser?", "Browsers cache timetable data for speed. The final staging-only version update tells them to fetch the refreshed data."),
        ("Can the staging Excel data feed be used?", "No. data-feed.html is intentionally excluded from Firebase Hosting to avoid an unnecessary export endpoint and extra reads in staging."),
        ("Where are credentials stored?", "Service-account keys are held by authorized operators outside the repository. Passwords and JSON keys should never be added to source control or shared in this guide."),
        ("Who should run a refresh?", "Only a trusted technical operator with the two purpose-limited Google Cloud service-account keys and the staging runbook."),
    ]:
        story.append(p(q, PDF_STYLES["question"]))
        story.append(p(a, PDF_STYLES["body"]))
    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont("Helvetica", 8); canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(0.72*inch, 0.37*inch, "UCVM Staging Environment Guide")
        canvas.drawRightString(7.78*inch, 0.37*inch, f"Page {doc.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    diagram()
    build_docx()
    build_pdf()
    print(DOCX_OUT)
    print(PDF_OUT)
